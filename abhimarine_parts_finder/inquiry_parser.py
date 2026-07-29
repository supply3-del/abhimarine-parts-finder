"""
Customer inquiry ingestion — turns a dumped Excel / Word / PDF / CSV / image
file (in whatever table layout the customer used) into a canonical DataFrame:
   item_no | description | part_number | qty

Column headers vary wildly between customers (SR NO, ITEM, DESCRIPTION,
PART NO/IMPA NO, QTY REQUIRED, Qty. Ordered, ...) so headers are matched by
keyword rather than fixed position, with a content-based fallback (regex /
numeric sniffing) for the cases where header text alone is ambiguous.

Parsing here is best-effort — the caller shows the result in an editable
grid so the team can fix any row before matching against inventory.
"""
import io
import re
from typing import List, Optional

import pandas as pd

# ----------------------------------------------------------------- schema
# field -> header keywords, most specific first (substring match, lowercase)
FIELD_SYNONYMS = {
    "part_number": ["part no./ref.no", "part no/impa", "part no. /ref", "ref.no", "ref no",
                    "part number", "part no", "drawing no", "drg no", "impa no"],
    "description": ["description", "part name", "particulars", "item description"],
    "qty": ["quantity required", "qty required", "qty reqd", "qty. ordered", "qty ordered",
            "quantity ordered", "quantity", "qty"],
    "item_no": ["sr no", "sr.no", "s.no", "sl no", "item no", "item", "no"],
    "brand": ["maker details", "maker detail", "brand name", "maker", "make", "brand"],
    "model": ["model/type", "model / type", "type / model", "type/model",
             "model no", "model number", "model", "type"],
}
CANON_COLS = ["item_no", "description", "part_number", "qty", "brand", "model"]

_PARTNO_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9\-/.]{3,}$")


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _header_map(headers: List[str]) -> dict:
    """Map column index -> canonical field, using header text keywords."""
    lower = [_norm(h).lower() for h in headers]
    used, mapping = set(), {}
    for field, synonyms in FIELD_SYNONYMS.items():
        for syn in synonyms:
            hit = next((i for i, h in enumerate(lower) if i not in used and syn in h), None)
            if hit is not None:
                mapping[field] = hit
                used.add(hit)
                break
    return mapping


def _looks_like_partno(s: str) -> bool:
    s = s.strip()
    if not s or s.upper() in ("NA", "N/A", "NIL", "-"):
        return False
    return bool(_PARTNO_RE.match(s)) and any(ch.isdigit() for ch in s)


def _fill_gaps_by_content(rows: List[List[str]], mapping: dict, ncols: int) -> dict:
    """For canonical fields the header pass couldn't place, guess from cell content."""
    used = set(mapping.values())
    free = [c for c in range(ncols) if c not in used]
    if not free or not rows:
        return mapping

    def col_vals(c):
        return [_norm(r[c]) for r in rows if c < len(r) and _norm(r[c])]

    def majority(vals, pred, thresh=0.6):
        return bool(vals) and sum(pred(v) for v in vals) >= max(1, len(vals) * thresh)

    is_int = lambda v: bool(re.match(r"^\d+(\.\d+)?$", v.replace(",", "")))

    # leftmost numeric column defaults to item_no (a leading serial number is near-universal
    # in these tables) so it isn't mistaken for qty just because both are plain digits
    if "item_no" not in mapping and free:
        c0 = free[0]
        if majority(col_vals(c0), is_int):
            mapping["item_no"] = c0; free.remove(c0)

    if "qty" not in mapping:
        for c in reversed(free):
            if majority(col_vals(c), is_int):
                mapping["qty"] = c; free.remove(c); break

    if "part_number" not in mapping:
        for c in free:
            if majority(col_vals(c), _looks_like_partno, thresh=0.5):
                mapping["part_number"] = c; free.remove(c); break

    if "description" not in mapping and free:
        # widest / most word-rich remaining column
        best = max(free, key=lambda c: sum(len(_norm(r[c]).split()) for r in rows if c < len(r)))
        mapping["description"] = best; free.remove(best)

    if "item_no" not in mapping and free:
        mapping["item_no"] = free[0]

    return mapping


def _rows_to_canonical(header: List[str], rows: List[List[str]],
                       mapping: Optional[dict] = None) -> pd.DataFrame:
    if mapping is None:
        mapping = _header_map(header)
        mapping = _fill_gaps_by_content(rows, mapping, len(header))

    out = []
    seen_numeric_item = False
    for i, r in enumerate(rows, 1):
        def get(field):
            c = mapping.get(field)
            return _norm(r[c]) if c is not None and c < len(r) else ""
        item_raw = get("item_no")
        desc = get("description")
        pno = get("part_number")
        # once real (numbered) rows are underway, a colon/sentence-like item_no (e.g.
        # "Chief Engineer:  ...", "Date:  ...") signals a trailing footer / signature
        # block rather than more table data — OCR noise alone shouldn't trigger this
        looks_like_footer = ":" in item_raw or len(item_raw.split()) >= 3
        if item_raw and seen_numeric_item and not re.match(r"^\d+$", item_raw) and looks_like_footer:
            break
        if re.match(r"^\d+$", item_raw):
            seen_numeric_item = True
        if not desc and not pno:
            continue
        qty_raw = get("qty").replace(",", "")
        m = re.match(r"(\d+(?:\.\d+)?)", qty_raw)
        if m:
            qty = float(m.group(1))
            qty = int(qty) if qty == int(qty) else qty
        else:
            qty = 1
        out.append({"item_no": get("item_no") or i, "description": desc,
                    "part_number": "" if pno.upper() in ("NA", "N/A", "NIL", "-") else pno,
                    "qty": qty, "brand": get("brand"), "model": get("model")})
    return pd.DataFrame(out, columns=CANON_COLS)


def _find_header_row(grid: List[List[str]], max_scan=40) -> Optional[int]:
    best_idx, best_score = None, 1  # need at least 2 field hits to count as a header
    for i, row in enumerate(grid[:max_scan]):
        mapping = _header_map(row)
        if len(mapping) > best_score:
            best_score, best_idx = len(mapping), i
    return best_idx


# requisition sheets carry an equipment header above the parts table, e.g.
#   MAKER'S FULL DETAILS: CATERPILLER   |   TYPE/MAKE/MODEL: TYPE- C18 S/N: CYG00156
# giving one brand+model for the whole sheet. The parts table itself rarely has
# per-row brand/model columns, so we read them from these label:value cells.
_EQUIP_SN_RE = re.compile(r"\bS\.?\s*/?\s*N\.?\b.*$", re.I)   # drop the serial-number tail


def _clean_equip_val(v: str) -> str:
    v = _EQUIP_SN_RE.sub("", v)
    v = re.sub(r"^\s*(type|make|model)\b[-:\s]*", "", v, flags=re.I)
    return re.sub(r"\s+", " ", v).strip(" -:/")


def _grid_equipment(header_rows: List[List[str]]):
    """Scan the cells above the parts table for a maker/make and model/type
    label:value, returning one (brand, model) for the whole sheet. The dedicated
    Maker field wins over Machinery/Equipment (which names the equipment type,
    not its maker) even though it usually appears lower on the sheet."""
    maker = equip = model = ""
    for row in header_rows:
        for cell in row:
            c = _norm(cell)
            if ":" not in c:
                continue
            label, _, val = c.partition(":")
            label, val = label.strip().lower(), val.strip()
            if not val:
                continue
            if not maker and re.search(r"maker|make|brand", label) and "model" not in label:
                maker = _clean_equip_val(val)
            elif not equip and re.search(r"machinery|equipment", label):
                equip = _clean_equip_val(val)
            if not model and re.search(r"model|type", label) and "requisition" not in label:
                model = _clean_equip_val(val)
    return (maker or equip), model


def _grid_to_canonical(grid: List[List[str]]) -> pd.DataFrame:
    grid = [r for r in grid if any(_norm(c) for c in r)]
    if not grid:
        return pd.DataFrame(columns=CANON_COLS)
    hdr_idx = _find_header_row(grid)
    if hdr_idx is None:
        return pd.DataFrame(columns=CANON_COLS)
    df = _rows_to_canonical(grid[hdr_idx], grid[hdr_idx + 1:])
    if not df.empty:
        brand, model = _grid_equipment(grid[:hdr_idx])
        if brand:
            df.loc[df["brand"] == "", "brand"] = brand
        if model:
            df.loc[df["model"] == "", "model"] = model
    return df


def _grid_to_canonical_by_content(grid: List[List[str]]) -> pd.DataFrame:
    """For OCR'd grids: header text is too unreliable to keyword-match, so drop the
    first (header) row and infer columns purely from what the data looks like."""
    grid = [r for r in grid if any(_norm(c) for c in r)]
    if len(grid) < 2:
        return pd.DataFrame(columns=CANON_COLS)
    header, rows = grid[0], grid[1:]
    mapping = _fill_gaps_by_content(rows, {}, len(header))
    return _rows_to_canonical(header, rows, mapping)


# ----------------------------------------------------------------- excel / csv
def parse_excel(data: bytes) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    frames = []
    for ws in wb.worksheets:
        grid = [[_norm(c) for c in row] for row in ws.iter_rows(values_only=True)]
        df = _grid_to_canonical(grid)
        if not df.empty:
            frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=CANON_COLS)


def parse_csv(data: bytes) -> pd.DataFrame:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="ignore")
    grid = [line.split(",") for line in text.splitlines()]
    return _grid_to_canonical(grid)


# a free-text part line, e.g. "Fuel Valve, Complet - 3 Pcs" (email/chat inquiries
# have no table structure at all, just one item per line ending in qty+unit)
_PROSE_LINE_RE = re.compile(
    r"^(?P<desc>.+?)\s*[–—-]\s*(?P<qty>\d+(?:\.\d+)?)\s*(?P<unit>pcs?|nos?|sets?|ea|no)\.?\s*$",
    re.I)


def _parse_prose(text: str) -> pd.DataFrame:
    """Fallback for pasted email/WhatsApp-style text with no table at all: part
    lines read '<description> - <qty> <unit>', and the equipment (Maker/Model/Type)
    is scattered across separate 'Label: value' lines rather than a header row."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    brand, model = _grid_equipment([[l] for l in lines])
    out = []
    for l in lines:
        m = _PROSE_LINE_RE.match(l)
        if not m:
            continue
        out.append({"item_no": len(out) + 1, "description": m["desc"].strip(" ,"),
                    "part_number": "", "qty": int(float(m["qty"])), "brand": brand, "model": model})
    return pd.DataFrame(out, columns=CANON_COLS)


def parse_pasted_text(text: str) -> pd.DataFrame:
    """Text pasted straight into the app — copied from Excel (tab-separated),
    a webpage table, selected/copied out of a PDF viewer (usually space-padded),
    or a plain email/chat inquiry with no table structure at all."""
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return pd.DataFrame(columns=CANON_COLS)

    sample = lines[:5]
    if sum("\t" in l for l in sample) >= max(1, len(sample) // 2):
        grid = [l.split("\t") for l in lines]
    elif sum(l.count(",") for l in sample) and all(l.count(",") >= 1 for l in sample):
        grid = [l.split(",") for l in lines]
    else:
        grid = [re.split(r"\s{2,}", l.strip()) for l in lines]

    grid = [[c.strip() for c in row] for row in grid]
    ncols = max(len(r) for r in grid)
    grid = [r + [""] * (ncols - len(r)) for r in grid]
    df = _grid_to_canonical(grid)
    return df if not df.empty else _parse_prose(text)


# ----------------------------------------------------------- text-layout requisitions
# Some requisition PDFs carry perfectly good text but no table grid that
# find_tables() can see — the columns are laid out by position, so the text
# extractor emits one cell per line. Those are parsed here, line-oriented.
_UOM = {"PCE", "PCS", "PC", "EA", "SET", "SETS", "NO", "NOS", "LTR", "KG", "MTR", "PAIR"}

# MAN B&W / MAK drawing numbers: 51610-11H-046, 61101-11 H, 50601-7H-068
_DRAWING_RE = re.compile(r"^(\d{4,6}-\d{1,3})\s?([A-Z])?$")
_FULL_PN_RE = re.compile(r"\b(\d{4,6}-\d{1,3}[A-Z]?-\d{2,4})\b")

_LABEL_BRAND_RE = re.compile(r"^(?:makers?|manufacturers?|make|brand)\s*:\s*(.*)$", re.I)
_LABEL_MODEL_RE = re.compile(r"^(?:machinery type|model|type)\s*[-:]\s*(.*)$", re.I)


def _qty_num(s: str) -> int:
    """'3,00' / '12.00' / '15PCS' -> int. Requisitions use the European decimal
    comma, so a bare ',00' tail is a decimal point, never a thousands group."""
    m = re.search(r"(\d+(?:[.,]\d+)?)", s or "")
    if not m:
        return 1
    return max(1, int(float(m.group(1).replace(",", "."))))


def _labelled_equipment(lines):
    """Brand/model from a requisition's header block. The label and its value are
    sometimes on the same line ('Makers: MAN B&W') and sometimes split across two
    ('Makers:' / 'MAN B&W'), so an empty capture takes the next non-blank line."""
    brand = model = ""
    for i, l in enumerate(lines):
        nxt = next((x for x in lines[i + 1:i + 3] if x.strip()), "")
        mb = _LABEL_BRAND_RE.match(l)
        if mb and not brand:
            brand = (mb.group(1).strip() or nxt).strip()
        mm = _LABEL_MODEL_RE.match(l)
        if mm and not model:
            model = (mm.group(1).strip() or nxt).strip()
        # scanned requisition header: 'CME MAN B&W ZJMD' / 'TYPE- 6L16/24'
        if not brand and l.upper().startswith("CME "):
            brand = l[4:].strip()
    # 'Model: X Serial #: Y' and 'Machinery type: X 1200RPM' run together
    model = re.split(r"\s+serial\s*#?\s*[:.]", model, flags=re.I)[0].strip()
    return brand, model


def _parse_shipserv(text: str) -> pd.DataFrame:
    """ShipServ/Wilhelmsen RFQ. Each 'Equipment Section Name:' block carries its
    own Manufacturer/Model, then items repeat as:
        <item no> / <part type> / <part number> / <description>
        [Buyer Comments: ... free text ...] / <UoM> / <qty>
    Brand and model are per-section, not per-document — one RFQ mixes makers."""
    lines = [l.strip() for l in text.splitlines()]
    brand = model = ""
    out, i = [], 0
    while i < len(lines):
        l = lines[i]
        m = re.match(r"Manufacturer:\s*(.*)", l, re.I)
        if m:
            # 'Manufacturer: MAK Model: 8M522' collapses onto one line sometimes
            split = re.search(r"(.*?)\s*Model:\s*(.*)", m.group(1), re.I)
            brand, model = ((split.group(1).strip(), split.group(2).strip())
                            if split else (m.group(1).strip(), model))
            model = re.split(r"\s+serial\s*#", model, flags=re.I)[0].strip()
            i += 1
            continue
        m = re.match(r"Model:\s*(.*)", l, re.I)
        if m:
            model = re.split(r"\s+serial\s*#", m.group(1), flags=re.I)[0].strip()
            i += 1
            continue
        if re.fullmatch(r"\d{1,3}", l) and i + 3 < len(lines):
            typ, pn, desc = lines[i + 1], lines[i + 2], lines[i + 3]
            if re.fullmatch(r"[A-Z]{2,3}", typ) and pn and desc:
                # the UoM/qty pair closes the block; Buyer Comments sit in between
                qty, j = None, i + 4
                while j < len(lines) and j < i + 25:
                    if lines[j].upper() in _UOM and re.fullmatch(r"\d+", lines[j + 1:j + 2] and lines[j + 1] or ""):
                        qty = int(lines[j + 1])
                        break
                    j += 1
                out.append({"item_no": l, "description": desc.strip(),
                           "part_number": pn.strip(), "qty": qty or 1,
                           "brand": brand, "model": model})
                i = (j + 2) if qty else (i + 4)
                continue
        i += 1
    return pd.DataFrame(out, columns=CANON_COLS)


def _parse_numbered_requisition(text: str) -> pd.DataFrame:
    """Scanned requisition print-outs (post-OCR), one item per line:
        '1. BALL BEARING 51610-11H-046 PC 3,00 0,00'
    The part number is embedded in the description text, and occasionally spills
    onto the following line ('... FOR HT COOLING WATER PC 4,00' / '51103-17H-278')."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    brand, model = _labelled_equipment(lines)
    row_re = re.compile(
        r"^(\d{1,3})\s*[.,)]\s+(.+?)\s+(" + "|".join(_UOM) + r")\s+([\d.,]+)\s+[\d.,]+\s*$", re.I)
    out = []
    for idx, l in enumerate(lines):
        m = row_re.match(l)
        if not m:
            continue
        item_no, middle, _uom, qty = m.groups()
        pn_hit = _FULL_PN_RE.search(middle)
        if pn_hit:
            pn = pn_hit.group(1)
            desc = (middle[:pn_hit.start()] + " " + middle[pn_hit.end():])
        else:
            # part number wrapped to the next line, bare or parenthesised
            nxt = lines[idx + 1] if idx + 1 < len(lines) else ""
            nxt_hit = _FULL_PN_RE.search(nxt) if not row_re.match(nxt) else None
            pn = nxt_hit.group(1) if nxt_hit else ""
            desc = middle
        out.append({"item_no": item_no, "description": re.sub(r"\s{2,}", " ", desc).strip(" ,.-"),
                   "part_number": pn, "qty": _qty_num(qty), "brand": brand, "model": model})
    return pd.DataFrame(out, columns=CANON_COLS)


def _parse_drawing_group_table(text: str) -> pd.DataFrame:
    """Requisitions that list a drawing-group header ('FRAME WITH MAIN BEARING'
    / '61101-11 H') and then number items by their position *within* that
    drawing — the part number cell holds only a suffix ('06', '13'). The real
    inventory part number is group+suffix ('61101-11H-06'), so compose it;
    a bare '06' is unmatchable on its own."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    brand, model = _labelled_equipment(lines)
    out, drawing, i = [], "", 0
    while i < len(lines):
        m = _DRAWING_RE.match(lines[i])
        if m:
            drawing = m.group(1) + (m.group(2) or "")   # '61101-11 H' -> '61101-11H'
            i += 1
            continue
        # item block: <item no> / <description> / <part no> / <qty + unit>
        if re.fullmatch(r"\d{1,3}", lines[i]) and i + 3 < len(lines):
            item_no, desc, pn, qty = lines[i:i + 4]
            if re.match(r"^\d+\s*(" + "|".join(_UOM) + r")\.?$", qty, re.I) and desc:
                if re.fullmatch(r"\d{1,3}", pn):
                    # A bare position number only means something against its
                    # drawing. Without one it is not a part number at all —
                    # emitting it raw matches junk rows whose part_number is
                    # literally '12', so drop it and match on description.
                    pn = f"{drawing}-{pn}" if drawing else ""
                out.append({"item_no": item_no, "description": desc, "part_number": pn,
                           "qty": _qty_num(qty), "brand": brand, "model": model})
                i += 4
                continue
        i += 1
    return pd.DataFrame(out, columns=CANON_COLS)


def _identified(df: pd.DataFrame) -> int:
    """How many rows came out with an actual part number."""
    if df is None or df.empty or "part_number" not in df:
        return 0
    return int((df["part_number"].astype(str).str.strip() != "").sum())


def best_parse(*frames: pd.DataFrame) -> pd.DataFrame:
    """Pick the parse that identified the most parts, falling back to row count.
    Raw row count alone is a bad judge: an OCR grid emits one row per letterhead
    line and will out-count a clean extraction 106 to 79. Ties go to the earliest
    argument, so callers pass their highest-fidelity parser first."""
    return max(frames, key=lambda d: (_identified(d), len(d) if d is not None else 0))


def parse_text_layout(text: str) -> pd.DataFrame:
    """Try each positional-text requisition layout, keep the best result. Cheaper
    and steadier than sniffing the format up front — these documents have no
    reliable marker, and a wrong guess loses the whole file."""
    frames = []
    for parser in (_parse_shipserv, _parse_numbered_requisition, _parse_drawing_group_table):
        try:
            frames.append(parser(text))
        except Exception:
            continue
    return best_parse(*frames) if frames else pd.DataFrame(columns=CANON_COLS)


# ----------------------------------------------------------------- word (.doc / .docx)
# Main/aux-engine header line, e.g. 'M/E MAK  6M 551AK SN:55346' — the maker and
# model apply to the whole requisition but sit in the header, not the parts table.
_ENGINE_HDR_RE = re.compile(
    r"\b[MA]\s*/\s*E\b\s*[:\-]?\s*(.+?)\s*(?:S\.?/?N|SERIAL|DEPARTMENT|IMO|$)", re.I)


def _word_header_equipment(text: str) -> str:
    """Maker+model descriptor from a ship requisition header — either the labelled
    form ('Maker:'/'Model:') or the 'M/E <maker> <model> SN:...' engine line. The
    raw descriptor is returned; annotate_brand_model normalises it to inventory."""
    lines = [l for l in re.split(r"[\r\n]+", text or "") if l.strip()]
    b, m = _labelled_equipment(lines)
    if b or m:
        return (b + " " + m).strip()
    hit = _ENGINE_HDR_RE.search(text or "")
    return re.sub(r"\s+", " ", hit.group(1)).strip() if hit else ""


def parse_word(path: str) -> pd.DataFrame:
    """Requires MS Word installed (uses COM automation) — works for legacy .doc and .docx."""
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    frames, header = [], ""
    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(path, False, True)
        try:
            header = _word_header_equipment(doc.Content.Text)
            for tbl in doc.Tables:
                nrows, ncols = tbl.Rows.Count, tbl.Columns.Count
                grid = []
                for r in range(1, nrows + 1):
                    row = []
                    for c in range(1, ncols + 1):
                        try:
                            cell = tbl.Cell(r, c).Range.Text
                        except Exception:
                            cell = ""
                        row.append(cell.replace("\x07", "").replace("\r", "").strip())
                    grid.append(row)
                df = _grid_to_canonical(grid)
                if not df.empty:
                    frames.append(df)
        finally:
            doc.Close(False)
    finally:
        word.Quit()
    out = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=CANON_COLS)
    # the header maker/model belongs to every row that didn't carry its own
    if header and not out.empty:
        for col in ("brand", "model"):
            blank = out[col].fillna("").astype(str).str.strip() == ""
            out.loc[blank, col] = header
    return out


# ----------------------------------------------------------------- word binary .doc (no Word/COM)
# Legacy binary .doc stores its text as UTF-16LE runs. When MS Word/COM isn't
# available (or its license has expired) we can still recover the requisition:
# the equipment header (Type/Make/Maker/Model) gives brand+model for the whole
# sheet, and the parts table is rebuilt by anchoring on each qty cell.
# ponytail: tuned to the ship-requisition layout we actually receive; a wildly
#   different .doc table may need COM (parse_word). Upgrade path = LibreOffice
#   headless convert-to-docx if these files diversify.
_DOC_NOISE = ("Normal", "Heading", "Default Paragraph", "Table ", "No List", "Header",
              "Footer", "Footnote", "Endnote", "Line Number", "Hyperlink", "Strong",
              "Emphasis", "Times New Roman", "Calibri", "Arial", "PMingLiU", "Grammarly",
              "Root Entry", "WordDocument", "SummaryInformation", "MsoDataStore",
              "Properties", "Normal Table")
_QTYUNIT_RE = re.compile(r"^\s*(\d+)\s*(pcs?|nos?|set|ea|no)\b", re.I)   # "30PCS", "12 PCS"
_DOC_NOISE_CELL_RE = re.compile(r"^\s*(nil|na|n/a|-|\d+(\.\d+)?)\s*$", re.I)  # NIL / bare number
_DOC_HDR_RE = re.compile(r"qty|part\s*no|date|rob|remarks|approved|ordered|recd|received|item", re.I)


def _doc_runs(data: bytes):
    text = data.decode("utf-16-le", "ignore")
    runs = [r.strip() for r in re.findall(r"[ -~]{3,}", text)]
    return [r for r in runs if r and not any(n in r for n in _DOC_NOISE)]


def parse_doc(data: bytes) -> pd.DataFrame:
    runs = _doc_runs(data)
    if not runs:
        return pd.DataFrame(columns=CANON_COLS)

    # equipment header: a label run followed by its value run
    def after(*labels):
        for i, r in enumerate(runs):
            if any(r.rstrip(":").strip().lower() == l for l in labels) and i + 1 < len(runs):
                return runs[i + 1]
        return ""
    make = after("make", "maker", "brand")
    model = after("model", "type", "model no")

    # table region: from the "Description" header run to the signature footer
    start = next((i for i, r in enumerate(runs) if "description" in r.lower()), None)
    if start is None:
        return pd.DataFrame(columns=CANON_COLS)
    end = next((i for i, r in enumerate(runs[start + 1:], start + 1)
                if re.search(r"chief engineer|approved by|signature:|requisitioned|captain|"
                             r"tech\.?\s*supdt|prepared by", r, re.I)), len(runs))
    body = runs[start + 1:end]
    while body and _DOC_HDR_RE.search(body[0]):   # skip the remaining column-header cells
        body.pop(0)

    # rebuild rows as a state machine: a description opens a row; the first
    # "N PCS" is its qty; the first part-number-shaped token is its part number.
    # everything else in the row (a second qty column, NIL, bare item numbers)
    # is ignored — that's what caused phantom rows before.
    out, cur = [], None
    for r in body:
        qm = _QTYUNIT_RE.match(r)
        if qm:
            if cur and cur["qty"] is None:
                cur["qty"] = int(qm.group(1))
            continue
        if _DOC_NOISE_CELL_RE.match(r):
            continue
        if _looks_like_partno(r):
            if cur and not cur["part_number"]:
                cur["part_number"] = r
            continue
        cur = {"item_no": len(out) + 1, "description": r, "part_number": "",
               "qty": None, "brand": make, "model": model}
        out.append(cur)
    for row in out:
        row["qty"] = 1 if row["qty"] is None else row["qty"]
    return pd.DataFrame(out, columns=CANON_COLS)


# ----------------------------------------------------------------- pdf
def parse_pdf(data: bytes) -> pd.DataFrame:
    import fitz  # PyMuPDF
    frames = []
    doc = fitz.open(stream=data, filetype="pdf")
    text_pages = []
    for page in doc:
        found_table = False
        try:
            tabs = page.find_tables()
            for tab in tabs.tables:
                grid = [[_norm(c) for c in row] for row in tab.extract()]
                df = _grid_to_canonical(grid)
                if not df.empty:
                    frames.append(df)
                    found_table = True
        except Exception:
            pass
        text = page.get_text()
        if len(text.strip()) >= 200:
            # Real text on the page. Even if no table was recognised, OCR'ing it
            # would only re-read the same characters worse — the columns are laid
            # out by position, so hand it to the text-layout parsers instead.
            text_pages.append(text)
        elif not found_table:
            # genuinely a scanned page — rasterize and OCR
            ocr = parse_image(page.get_pixmap(dpi=300).tobytes("png"))
            if not ocr.empty:
                frames.append(ocr)
            text_pages.append(_ocr_text(page))

    # The positional layouts (per-section brand/model, drawing-group part number
    # prefixes) only resolve with the whole document in view, so they run once
    # over the joined text rather than per page.
    whole = parse_text_layout("\n".join(text_pages))
    table_rows = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame(columns=CANON_COLS)
    # text-layout first: on an equal count it carries the better brand/model,
    # because it reads the per-section labels the table grid flattens away
    return best_parse(whole, table_rows)


def _ocr_text(page) -> str:
    """Raw OCR text of a scanned page, for the text-layout parsers (parse_image
    reconstructs a grid, which loses the line structure they key off)."""
    import pytesseract
    from PIL import Image
    _tesseract_path()
    img, _ = _preprocess_for_ocr(
        Image.open(io.BytesIO(page.get_pixmap(dpi=300).tobytes("png"))).convert("RGB"))
    return pytesseract.image_to_string(img, config="--psm 6")


# ----------------------------------------------------------------- image (OCR)
def _preprocess_for_ocr(img):
    """Grayscale + upscale + threshold — small/low-res phone screenshots of tables
    OCR far more reliably after this than as-is."""
    from PIL import Image, ImageOps
    img = img.convert("L")
    scale = 3 if max(img.size) < 1500 else 1
    if scale > 1:
        img = img.resize((img.width * scale, img.height * scale), Image.LANCZOS)
    img = ImageOps.autocontrast(img)
    img = img.point(lambda p: 255 if p > 150 else 0)
    return img, scale


def _tesseract_path():
    """Point pytesseract at the Windows install when it isn't on PATH."""
    import os
    import pytesseract
    if not pytesseract.pytesseract.tesseract_cmd or pytesseract.pytesseract.tesseract_cmd == "tesseract":
        for cand in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
            if os.path.exists(cand):
                pytesseract.pytesseract.tesseract_cmd = cand
                break


def parse_image(data: bytes) -> pd.DataFrame:
    """Best-effort OCR table reconstruction for photographed/scanned inquiry sheets."""
    import pytesseract
    from PIL import Image

    _tesseract_path()
    img = Image.open(io.BytesIO(data)).convert("RGB")
    img, _scale = _preprocess_for_ocr(img)
    text = pytesseract.image_to_string(img, config="--psm 6")
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return pd.DataFrame(columns=CANON_COLS)

    # bordered tables get their vertical rules transcribed as literal "|" characters —
    # that's a far more reliable column delimiter than guessing from word spacing.
    pipe_lines = sum(1 for l in lines if l.count("|") >= 2)
    use_pipes = pipe_lines >= max(2, len(lines) // 3)

    grid = []
    for l in lines:
        if use_pipes:
            cells = [c.strip(" |") for c in l.split("|")]
        else:
            cells = [c.strip() for c in re.split(r"\s{2,}", l.strip())]
        grid.append(cells)

    ncols = max(len(r) for r in grid)
    grid = [r + [""] * (ncols - len(r)) for r in grid]
    return _grid_to_canonical_by_content(grid)


# ------------------------------------------------ legacy .doc without MS Word
def _parse_doc_text(text: str) -> pd.DataFrame:
    """Parse the plain text of a .doc into rows — a space-delimited grid and the
    positional text-layout parsers, best result wins — then stamp the header
    maker/model onto any row that didn't carry its own."""
    header = _word_header_equipment(text)
    lines = [l for l in (text or "").splitlines() if l.strip()]
    grid = [re.split(r"\s{2,}", l.strip()) for l in lines]
    df_grid = pd.DataFrame(columns=CANON_COLS)
    if grid:
        ncols = max(len(r) for r in grid)
        grid = [r + [""] * (ncols - len(r)) for r in grid]
        if ncols >= 2:
            df_grid = _grid_to_canonical(grid)
    out = best_parse(df_grid, parse_text_layout(text))
    if header and not out.empty:
        for col in ("brand", "model"):
            blank = out[col].fillna("").astype(str).str.strip() == ""
            out.loc[blank, col] = header
    return out


def _parse_doc_antiword(path: str) -> pd.DataFrame:
    """Linux fallback for legacy binary .doc when MS Word/COM isn't available
    (e.g. Streamlit Cloud): shell out to antiword — a tiny apt package listed in
    packages.txt — for the document text, then parse it. Returns empty if
    antiword isn't installed, so the caller can fall back further."""
    import shutil
    import subprocess
    exe = shutil.which("antiword")
    if not exe:
        return pd.DataFrame(columns=CANON_COLS)
    try:
        res = subprocess.run([exe, "-w", "0", path], capture_output=True, timeout=30)
        text = res.stdout.decode("utf-8", "replace")
    except Exception:
        return pd.DataFrame(columns=CANON_COLS)
    return _parse_doc_text(text)


# ----------------------------------------------------------------- dispatch
def parse_inquiry(filename: str, data: bytes) -> pd.DataFrame:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm", "xls"):
        return parse_excel(data)
    if ext == "csv":
        return parse_csv(data)
    if ext in ("doc", "docx"):
        import tempfile, os
        with tempfile.NamedTemporaryFile(suffix="." + ext, delete=False) as f:
            f.write(data)
            tmp_path = f.name
        try:
            try:
                df = parse_word(tmp_path)        # MS Word COM (Windows)
            except Exception:
                df = pd.DataFrame(columns=CANON_COLS)
            if not df.empty:
                return df
            # Word/COM gave nothing (not installed / Linux host). For a binary
            # .doc, try antiword, then the pure-Python byte reader. .docx is a
            # zip the byte reader can't touch, so it just returns what we have.
            if ext == "doc":
                aw = _parse_doc_antiword(tmp_path)
                return aw if not aw.empty else parse_doc(data)
            return df
        finally:
            os.unlink(tmp_path)
    if ext == "pdf":
        return parse_pdf(data)
    if ext in ("jpg", "jpeg", "png", "bmp", "tiff"):
        return parse_image(data)
    raise ValueError(f"Unsupported file type: .{ext}")


if __name__ == "__main__":
    # binary .doc reconstruction: header brand/model + qty-anchored rows
    cells = ["Type", "WEICHAI X6170ZC", "Make", "WEICHAI 170Z SERIES",
             "Maker", "s details", "WEIFANG DIESEL", "Model", "WEICHAI 170Z X6170ZC",
             "Item", "Description", "Qty. Ordered", "Part No./Ref.No./Remarks", "ROB",
             "GASKET", "30PCS", "170Z.21.11/POS.NO4", "NIL",
             "B1 ORING42.5X3.75", "80PCS", "GB/T345.1", "NIL",
             "Chief Engineer:   SOMEONE", "TO BE FILLED BY REQUISITIONER"]
    fixture = "\x00".join(cells).encode("utf-16-le")
    df = parse_doc(fixture)
    assert len(df) == 2, df
    assert list(df["description"]) == ["GASKET", "B1 ORING42.5X3.75"], list(df["description"])
    assert list(df["part_number"]) == ["170Z.21.11/POS.NO4", "GB/T345.1"], list(df["part_number"])
    assert list(df["qty"]) == [30, 80], list(df["qty"])
    assert (df["brand"] == "WEICHAI 170Z SERIES").all()
    assert (df["model"] == "WEICHAI X6170ZC").all()

    # legacy .doc read as plain text (what antiword emits on the Linux Cloud host):
    # space-delimited columns + an 'M/E <maker> <model>' engine header
    doc_text = ("M/V : MED SEA   M/E MAK  6M 551AK  SN:55346\n"
                "No   Description       Part No     Qty   Unit\n"
                "01   INTEK VALVE       7.2210-2    2     PCS\n"
                "02   EXHUST VALVE      7.2210-2    4     PCS\n"
                "03   VALVE SEAT RING   7.2220-5    6     PCS\n")
    dt = _parse_doc_text(doc_text)
    assert len(dt) == 3, dt
    assert list(dt["part_number"]) == ["7.2210-2", "7.2210-2", "7.2220-5"], list(dt["part_number"])
    assert list(dt["qty"]) == [2, 4, 6], list(dt["qty"])
    assert (dt["brand"] == "MAK 6M 551AK").all(), list(dt["brand"])   # header stamped on every row
    # no antiword binary installed -> empty, so the caller falls back
    import shutil as _sh
    if _sh.which("antiword") is None:
        assert _parse_doc_antiword("nonexistent.doc").empty

    # free-text email/chat paste: no table, just prose lines + scattered label:value
    email = """We kindly invite your good company to send your quotation for the below request.

Fuel Valve, Complet –  3 Pcs
Set Parts For Fuel Valve – 3 Pcs
Set O-ring for Fuel Valve – 3 Pcs

Main Engine

Type Model: 7RTA84T

Maker: Diesel United-Sulzer

Serial No. 3126"""
    pdf = parse_pasted_text(email)
    assert list(pdf["description"]) == ["Fuel Valve, Complet", "Set Parts For Fuel Valve",
                                        "Set O-ring for Fuel Valve"], list(pdf["description"])
    assert list(pdf["qty"]) == [3, 3, 3], list(pdf["qty"])
    assert (pdf["brand"] == "Diesel United-Sulzer").all()
    assert (pdf["model"] == "7RTA84T").all()

    # ShipServ RFQ: per-section Manufacturer/Model, Buyer Comments between
    # the description and the UoM/qty pair, and a second section that switches maker
    shipserv = """Equipment Section Name: Diesel Engine No. 5
Manufacturer: MAK
Model: 8M552 Serial #: 57073
#
Part
Type
Part Number
Description
UoM
Qty
1
MF
7.3191-100
Camshaft Bearing, Compl.
Buyer Comments: [MAKER] MAK
[MTRL] Drawing:B1.05.03 Item No. 1.0-4
PCE
9
2
MF
7.1130-16
Gasket
PCE
16
Equipment Section Name: Diesel Generator No 5
Manufacturer: Siemens
Model: 1DK4933-8AL07-Z Serial #: 178537
3
BP
1.0520-096
Distance piece
PCE
8"""
    s = _parse_shipserv(shipserv)
    assert len(s) == 3, s
    assert list(s["part_number"]) == ["7.3191-100", "7.1130-16", "1.0520-096"], list(s["part_number"])
    assert list(s["qty"]) == [9, 16, 8], list(s["qty"])
    # brand/model follow the section, and 'Serial #' is stripped off the model
    assert list(s["brand"]) == ["MAK", "MAK", "Siemens"], list(s["brand"])
    assert list(s["model"]) == ["8M552", "8M552", "1DK4933-8AL07-Z"], list(s["model"])
    # the Buyer Comments block must not leak into the description
    assert s["description"].iloc[0] == "Camshaft Bearing, Compl.", s["description"].iloc[0]

    # scanned requisition (post-OCR): part number inline, or wrapped to the next
    # line, or parenthesised at the end of a wrapped description
    scanned = """CME MAN B&W ZJMD
TYPE- 6L16/24
1. BALL BEARING 51610-11H-046 PC 3,00 0,00
2. THERMOSTATIC ELEMENT FOR HT COOLING WATER PC 4,00 0,00
51103-17H-278
10. CONNECTING SOCKET COMPLETE INCL. VALVE CONES, PC 6,00 0,00
SPRINGS AND ITEM 577 (51401-07H-230)"""
    n = _parse_numbered_requisition(scanned)
    assert len(n) == 3, n
    assert list(n["part_number"]) == ["51610-11H-046", "51103-17H-278", "51401-07H-230"], \
        list(n["part_number"])
    assert list(n["qty"]) == [3, 4, 6], list(n["qty"])   # '3,00' is 3, not 300
    assert (n["brand"] == "MAN B&W ZJMD").all()
    assert (n["model"] == "6L16/24").all()

    # drawing-group table: items carry only their position within the drawing,
    # so '06' under '61101-11 H' has to be composed into '61101-11H-06'
    grouped = """Makers:
HYAUNDAI MAN B$W
Machinery type:
5L28/32H
FRAME WITH MAIN BEARING
61101-11 H
4
O-RING
06
15PCS
5
MAIN BEARING 2/2
13
7 PCS"""
    g = _parse_drawing_group_table(grouped)
    assert len(g) == 2, g
    assert list(g["part_number"]) == ["61101-11H-06", "61101-11H-13"], list(g["part_number"])
    assert list(g["qty"]) == [15, 7], list(g["qty"])
    assert (g["model"] == "5L28/32H").all()          # label and value on separate lines

    # arbitration: a junk-heavy OCR grid must not beat a clean identified parse
    clean = pd.DataFrame({"item_no": [1, 2], "description": ["A", "B"],
                         "part_number": ["X-1", "X-2"], "qty": [1, 1],
                         "brand": ["", ""], "model": ["", ""]})
    junk = pd.DataFrame({"item_no": range(9), "description": ["letterhead"] * 9,
                        "part_number": [""] * 9, "qty": [1] * 9,
                        "brand": [""] * 9, "model": [""] * 9})
    assert best_parse(clean, junk) is clean
    assert len(best_parse(junk, pd.DataFrame(columns=CANON_COLS))) == 9

    print("inquiry_parser self-check OK")
